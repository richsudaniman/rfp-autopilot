from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from rfp_autopilot import ExtractiveDrafter, LLMDrafter, Retriever, load_library, process_workbook
from rfp_autopilot.drafter import INSUFFICIENT
from rfp_autopilot.pipeline import STATUS_DRAFTED, STATUS_SME, Thresholds, answer_question

ROOT = Path(__file__).resolve().parents[1]
LIBRARY = load_library(ROOT / "data" / "answer_library.csv")


@pytest.fixture(scope="module")
def retriever():
    return Retriever(LIBRARY)


# ---------- library ----------

def test_library_loads_with_aliases():
    sso = next(e for e in LIBRARY if e.id == "IAM-001")
    assert "SAML" in sso.question
    ops = next(e for e in LIBRARY if e.id == "OPS-004")
    assert len(ops.phrasings) > 1


def test_library_rejects_missing_columns(tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text("id,question\n1,Do you encrypt?\n")
    with pytest.raises(ValueError, match="missing required columns"):
        load_library(bad)


def test_library_rejects_duplicate_ids(tmp_path):
    dup = tmp_path / "dup.csv"
    dup.write_text("id,question,answer\nA,Q1,A1\nA,Q2,A2\n")
    with pytest.raises(ValueError, match="duplicate"):
        load_library(dup)


# ---------- retrieval ----------

@pytest.mark.parametrize("question, expected_id", [
    ("Do you support single sign-on through SAML?", "IAM-001"),
    ("Can admins enforce 2FA?", "IAM-002"),
    ("Is data encrypted at rest?", "SEC-001"),
    ("What are your RTO and RPO?", "OPS-002"),
    ("Do you perform annual penetration testing by a third party?", "SEC-004"),
])
def test_retriever_finds_reworded_questions(retriever, question, expected_id):
    assert retriever.search(question)[0].entry.id == expected_id


def test_retriever_returns_k_distinct_entries(retriever):
    matches = retriever.search("encryption", k=3)
    assert len(matches) == 3
    assert len({m.entry.id for m in matches}) == 3
    assert matches[0].score >= matches[1].score >= matches[2].score


# ---------- guardrail ----------

def test_unknown_question_is_flagged_not_invented(retriever):
    answer, confidence, _, _, status = answer_question(
        "Are you FedRAMP authorized?", retriever, ExtractiveDrafter(), Thresholds())
    assert answer is None
    assert status == STATUS_SME
    assert confidence == "Low"


class FakeClient:
    """Stands in for anthropic.Anthropic(); records the prompt it was sent."""

    def __init__(self, reply):
        self.reply, self.calls = reply, []
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=self.reply)])


def test_llm_saying_insufficient_context_gets_flagged(retriever):
    client = FakeClient(INSUFFICIENT)
    answer, _, _, ids, status = answer_question(
        "Do you support SAML SSO with Okta and also OIDC?", retriever, LLMDrafter(client=client), Thresholds())
    assert answer is None and status == STATUS_SME
    assert "IAM-001" in ids


def test_llm_is_only_given_approved_answers(retriever):
    client = FakeClient("Yes, we support SAML 2.0 SSO with Okta.")
    answer, confidence, _, ids, status = answer_question(
        "Does your platform support SAML SSO with Okta?", retriever, LLMDrafter(client=client), Thresholds())
    assert answer.startswith("Yes") and status == STATUS_DRAFTED and confidence == "High"
    prompt = client.calls[0]["messages"][0]["content"]
    assert "[IAM-001]" in prompt and ids[0] == "IAM-001"
    assert "ONLY" in client.calls[0]["system"]


# ---------- end to end ----------

def _questionnaire(path, questions):
    wb = Workbook()
    ws = wb.active
    ws.append(["#", "Question"])
    for i, q in enumerate(questions, 1):
        ws.append([i, q])
    wb.save(path)


def test_workbook_round_trip(tmp_path, retriever):
    src, out = tmp_path / "q.xlsx", tmp_path / "q_answered.xlsx"
    _questionnaire(src, ["Do you support SAML SSO?", "Are you FedRAMP authorized?", "", "Is data encrypted in transit?"])

    report = process_workbook(src, out, retriever, ExtractiveDrafter())

    assert len(report.results) == 3  # blank row skipped
    ws = load_workbook(out).active
    headers = [c.value for c in ws[1]]
    assert headers[2:] == ["Draft Answer", "Confidence", "Match Score", "Source IDs", "Status"]
    assert ws["C2"].value.startswith("Yes") and ws["G2"].value == STATUS_DRAFTED
    assert ws["C3"].value in (None, "") and ws["G3"].value == STATUS_SME


def test_rerunning_reuses_output_columns(tmp_path, retriever):
    src, out = tmp_path / "q.xlsx", tmp_path / "q_answered.xlsx"
    _questionnaire(src, ["Do you support SAML SSO?"])
    process_workbook(src, out, retriever, ExtractiveDrafter())
    process_workbook(out, out, retriever, ExtractiveDrafter())
    assert load_workbook(out).active.max_column == 7

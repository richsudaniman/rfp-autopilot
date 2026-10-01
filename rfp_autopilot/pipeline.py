"""questionnaire.xlsx -> retrieve -> draft or flag -> answered.xlsx"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .drafter import Drafter
from .retriever import Retriever

OUTPUT_COLUMNS = ["Draft Answer", "Confidence", "Match Score", "Source IDs", "Status"]

STATUS_DRAFTED = "Drafted"
STATUS_REVIEW = "Drafted - quick review"
STATUS_SME = "Needs SME review"

FILLS = {
    STATUS_DRAFTED: PatternFill("solid", fgColor="E2F0D9"),  # green
    STATUS_REVIEW: PatternFill("solid", fgColor="FFF2CC"),   # yellow
    STATUS_SME: PatternFill("solid", fgColor="F8CBAD"),      # red
}


@dataclass
class Thresholds:
    min_match: float = 0.48   # below this, never draft: send to a human
    high: float = 0.65        # at or above this, confidence = High
    context_floor: float = 0.20  # matches below this aren't shown to the LLM


@dataclass
class RowResult:
    row: int
    question: str
    answer: str | None
    confidence: str
    score: float
    source_ids: list[str]
    status: str


@dataclass
class Report:
    results: list[RowResult] = field(default_factory=list)

    @property
    def counts(self) -> Counter:
        return Counter(r.status for r in self.results)

    def summary(self) -> str:
        c = self.counts
        total = len(self.results)
        return (
            f"{total} questions: {c[STATUS_DRAFTED]} drafted, "
            f"{c[STATUS_REVIEW]} need a quick review, {c[STATUS_SME]} flagged for SME review"
        )


def answer_question(question: str, retriever: Retriever, drafter: Drafter,
                    thresholds: Thresholds, k: int = 3) -> tuple[str | None, str, float, list[str], str]:
    matches = retriever.search(question, k=k)
    best = matches[0].score if matches else 0.0
    if best < thresholds.min_match:
        # Guardrail: nothing similar has been approved before. Don't guess.
        return None, "Low", best, [m.entry.id for m in matches[:1]], STATUS_SME

    context = [m for m in matches if m.score >= thresholds.context_floor]
    draft = drafter.draft(question, context)
    if draft.text is None:
        return None, "Low", best, draft.used_ids, STATUS_SME
    if best >= thresholds.high:
        return draft.text, "High", best, draft.used_ids, STATUS_DRAFTED
    return draft.text, "Medium", best, draft.used_ids, STATUS_REVIEW


def _find_question_column(ws, header_row: int) -> int:
    for cell in ws[header_row]:
        if isinstance(cell.value, str) and "question" in cell.value.lower():
            return cell.column
    raise ValueError(f"No column with 'question' in its header on row {header_row} of sheet '{ws.title}'")


def process_workbook(input_path: str | Path, output_path: str | Path, retriever: Retriever,
                     drafter: Drafter, thresholds: Thresholds | None = None,
                     sheet: str | None = None, header_row: int = 1) -> Report:
    thresholds = thresholds or Thresholds()
    wb = load_workbook(input_path)
    ws = wb[sheet] if sheet else wb.active
    q_col = _find_question_column(ws, header_row)

    # Reuse output columns if the sheet was processed before, else append them.
    headers = {str(c.value).strip(): c.column for c in ws[header_row] if c.value}
    next_col = ws.max_column + 1
    out_cols = {}
    for name in OUTPUT_COLUMNS:
        if name in headers:
            out_cols[name] = headers[name]
        else:
            out_cols[name] = next_col
            next_col += 1
        cell = ws.cell(row=header_row, column=out_cols[name], value=name)
        cell.font = Font(bold=True)

    report = Report()
    for row in range(header_row + 1, ws.max_row + 1):
        question = ws.cell(row=row, column=q_col).value
        if not isinstance(question, str) or not question.strip():
            continue
        answer, confidence, score, ids, status = answer_question(question, retriever, drafter, thresholds)
        values = {
            "Draft Answer": answer or "",
            "Confidence": confidence,
            "Match Score": round(score, 2),
            "Source IDs": ", ".join(ids),
            "Status": status,
        }
        for name, value in values.items():
            cell = ws.cell(row=row, column=out_cols[name], value=value)
            cell.fill = FILLS[status]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        report.results.append(RowResult(row, question, answer, confidence, score, ids, status))

    ws.column_dimensions[ws.cell(row=header_row, column=out_cols["Draft Answer"]).column_letter].width = 70
    ws.column_dimensions[ws.cell(row=header_row, column=out_cols["Status"]).column_letter].width = 22
    wb.save(output_path)
    return report
